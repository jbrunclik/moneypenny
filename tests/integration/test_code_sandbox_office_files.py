"""Live test: the sandbox image can produce Word, PowerPoint and Excel files.

The prompt tells the model the environment is fixed and not to probe it, so
a library the image lacks turns into a failed turn. Runs only where Docker
and the sandbox image are available (skips otherwise, like the isolation test).
"""

import json

import pytest

from tests.integration.test_code_sandbox_isolation import _sandbox_runnable

_OFFICE_SCRIPT = """
import docx, pptx, openpyxl
d = docx.Document()
d.add_heading('Trip plan', level=1)
d.add_paragraph('Day one', style='List Bullet')
d.save('/output/plan.docx')
p = pptx.Presentation()
s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = 'Trip plan'
p.save('/output/plan.pptx')
wb = openpyxl.Workbook()
wb.active['A1'] = 'Cost'
wb.active['A2'] = '=SUM(B1:B9)'
wb.save('/output/plan.xlsx')
print('OFFICE_OK')
"""


@pytest.mark.skipif(not _sandbox_runnable(), reason="Docker or sandbox image unavailable")
def test_sandbox_writes_office_files() -> None:
    from src.agent.tools.code_execution import execute_code

    result = json.loads(execute_code.invoke({"code": _OFFICE_SCRIPT}))

    assert result.get("success") is True, f"sandbox execution failed: {result}"
    assert "OFFICE_OK" in result["stdout"]
    names = {f["name"] for f in result["files"]}
    assert names == {"plan.docx", "plan.pptx", "plan.xlsx"}
